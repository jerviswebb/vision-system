import json
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

FOLDER = Path("/workspaces/vision-system/data/dashboard-demo")
REPORT = FOLDER / "weld_inspection_report.xlsx"

json_files = sorted(FOLDER.glob("*.json"))

if not json_files:
    raise RuntimeError(f"No JSON files found in {FOLDER}")

workbook = Workbook()
all_sheet = workbook.active
all_sheet.title = "All Inspections"
failed_sheet = workbook.create_sheet("Failed Inspections")

headers = [
    "Inspection ID",
    "Inspection Time (UTC)",
    "Result",
    "Selected Class",
    "Selected Confidence",
    "Model Profile",
    "Source Image",
    "Detection Details",
    "Annotated Image Path",
    "Source JSON",
]

for sheet in [all_sheet, failed_sheet]:
    sheet.append(headers)

for json_file in json_files:
    with json_file.open() as file:
        data = json.load(file)

    message = data.get("message", "")

    # Extract the image name from your existing message format.
    source_image = ""
    if message.startswith("Dataset demo: "):
        source_image = message.removeprefix("Dataset demo: ").split("; detections=", 1)[
            0
        ]

    # Excel stores dates without timezone information.
    # Convert to UTC first, then remove the timezone.
    timestamp = data.get("observedAt")
    inspection_time = None
    if timestamp:
        inspection_time = (
            datetime.fromisoformat(timestamp.replace("Z", "+00:00"))
            .astimezone(timezone.utc)
            .replace(tzinfo=None)
        )

    annotated_image = json_file.with_suffix(".jpg")

    row = [
        data.get("inspectionId"),
        inspection_time,
        data.get("state"),
        data.get("className"),
        data.get("confidence"),
        data.get("profile"),
        source_image,
        message,
        str(annotated_image) if annotated_image.exists() else "",
        json_file.name,
    ]

    all_sheet.append(row)

    if data.get("state") == "FAIL":
        failed_sheet.append(row)

# Format both sheets.
widths = [38, 24, 14, 20, 24, 20, 45, 65, 55, 45]

for sheet in [all_sheet, failed_sheet]:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="50439B")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(wrap_text=True)

    sheet.row_dimensions[1].height = 32

    for column, width in zip("ABCDEFGHIJ", widths):
        sheet.column_dimensions[column].width = width

    for row in sheet.iter_rows(min_row=2):
        row[1].number_format = "yyyy-mm-dd hh:mm:ss"
        row[4].number_format = "0.0%"
        row[7].alignment = Alignment(wrap_text=True, vertical="top")

        if row[2].value == "FAIL":
            row[2].fill = PatternFill("solid", fgColor="FCE4D6")

workbook.save(REPORT)

print(f"Excel report created: {REPORT}")
print(f"All inspections: {all_sheet.max_row - 1}")
print(f"Failed inspections: {failed_sheet.max_row - 1}")
